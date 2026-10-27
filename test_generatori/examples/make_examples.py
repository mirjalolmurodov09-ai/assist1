"""Namuna fayllarni yaratadi: examples/sample.docx va examples/sample.xlsx (50 ta savol).

Ishga tushirish:  python examples/make_examples.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook

# (savol, to'g'ri javob, [3 ta noto'g'ri javob])
BANK: list[tuple[str, str, list[str]]] = [
    ("Oʻzbekiston poytaxti qaysi shahar?", "Toshkent", ["Samarqand", "Buxoro", "Xiva"]),
    ("Oʻzbekiston Respublikasi mustaqilligi qachon e'lon qilingan?", "1991-yil 1-sentabr", ["1990-yil 20-iyun", "1992-yil 8-dekabr", "1989-yil 21-oktabr"]),
    ("Amudaryo qaysi dengizga quyilgan edi?", "Orol dengizi", ["Kaspiy dengizi", "Qora dengiz", "Baltika dengizi"]),
    ("Alisher Navoiy qaysi asrda yashagan?", "XV asr", ["X asr", "XIII asr", "XVII asr"]),
    ("«Oʻtkan kunlar» romanining muallifi kim?", "Abdulla Qodiriy", ["Choʻlpon", "Hamza", "Gʻafur Gʻulom"]),
    ("Sakkizga yettini koʻpaytirsak nechchi boʻladi?", "56", ["54", "48", "64"]),
    ("Toʻrtburchakning barcha burchaklari yigʻindisi necha gradus?", "360", ["180", "270", "720"]),
    ("Yer sayyorasining sun'iy yoʻldoshi qaysi?", "Oy", ["Mars", "Venera", "Quyosh"]),
    ("Suvning qaynash harorati (normal bosimda) necha °C?", "100", ["90", "80", "120"]),
    ("Eng katta okean qaysi?", "Tinch okeani", ["Atlantika okeani", "Hind okeani", "Shimoliy Muz okeani"]),
    ("Qaysi qit'ada Oʻzbekiston joylashgan?", "Osiyo", ["Yevropa", "Afrika", "Avstraliya"]),
    ("Amir Temur poytaxti qaysi shahar edi?", "Samarqand", ["Buxoro", "Xiva", "Termiz"]),
    ("Sinus 30° ning qiymati nechaga teng?", "1/2", ["√3/2", "√2/2", "1"]),
    ("Fotosintez jarayonida oʻsimlik qaysi gazni yutadi?", "Karbonat angidrid", ["Kislorod", "Azot", "Vodorod"]),
    ("Inson yuragida nechta kamera bor?", "4 ta", ["2 ta", "3 ta", "5 ta"]),
    ("Suvning kimyoviy formulasi qaysi?", "H2O", ["CO2", "NaCl", "O2"]),
    ("Yilda necha oy bor?", "12", ["10", "11", "13"]),
    ("Haftada necha kun bor?", "7", ["5", "6", "8"]),
    ("Eng uzun daryo Oʻzbekistonda qaysi?", "Amudaryo", ["Sirdaryo", "Zarafshon", "Chirchiq"]),
    ("Mirzo Ulugʻbek nima bilan shugʻullangan?", "Astronomiya", ["Tibbiyot", "Me'morchilik", "Dengizchilik"]),
    ("Ibn Sino qaysi sohada mashhur?", "Tibbiyot", ["Astronomiya", "Tarix", "Geografiya"]),
    ("Al-Xorazmiy qaysi fanning asoschilaridan biri?", "Algebra", ["Biologiya", "Geologiya", "Kimyo"]),
    ("Toʻgʻri burchakli uchburchakda gipotenuzaga nisbatan teorema qanday ataladi?", "Pifagor teoremasi", ["Viet teoremasi", "Talis teoremasi", "Nyuton teoremasi"]),
    ("100 ning 10 foizi nechaga teng?", "10", ["1", "20", "100"]),
    ("Kvadratning tomoni 5 sm boʻlsa, uning yuzi necha sm²?", "25", ["10", "20", "15"]),
    ("Quyidagilardan qaysi biri tub son?", "17", ["15", "21", "27"]),
    ("Ingliz tilida «kitob» soʻzi qanday?", "book", ["table", "pen", "door"]),
    ("«Qalam» soʻzi qaysi soʻz turkumiga kiradi?", "Ot", ["Sifat", "Fe'l", "Ravish"]),
    ("Oʻzbek alifbosida nechta harf bor (lotin yozuvida)?", "29", ["26", "32", "35"]),
    ("Oltin kimyoviy belgisi qaysi?", "Au", ["Ag", "Fe", "Cu"]),
    ("Havoning asosiy tarkibiy qismi qaysi gaz?", "Azot", ["Kislorod", "Argon", "Karbonat angidrid"]),
    ("Quyosh sistemasida eng katta sayyora qaysi?", "Yupiter", ["Saturn", "Yer", "Neptun"]),
    ("Tovush eng tez qaysi muhitda tarqaladi?", "Qattiq jismda", ["Havoda", "Vakuumda", "Bo'sh kosmosda"]),
    ("Elektr tokining kuchi qaysi birlikda oʻlchanadi?", "Amper", ["Volt", "Om", "Vatt"]),
    ("Yer yuzining qancha qismini suv egallagan (taxminan)?", "71%", ["50%", "30%", "90%"]),
    ("Qaysi hayvon sutemizuvchi?", "Delfin", ["Akula", "Losos baliq", "Qurbaqa"]),
    ("Qaysi shahar «Sharq gavhari» deb ataladi?", "Samarqand", ["Navoiy", "Andijon", "Nukus"]),
    ("Fargʻona vodiysi qaysi tog' tizmalari orasida joylashgan?", "Tyan-Shan va Pomir-Oloy", ["Himolay va Karakorum", "Alp va Pireney", "Ural va Kavkaz"]),
    ("Oʻzbekiston Respublikasi Konstitutsiyasi qachon qabul qilingan?", "1992-yil 8-dekabr", ["1991-yil 1-sentabr", "1995-yil 1-yanvar", "1990-yil 24-mart"]),
    ("Kompyuterning «miyasi» deb qaysi qurilma ataladi?", "Protsessor", ["Monitor", "Sichqoncha", "Klaviatura"]),
    ("1 kilobayt necha baytga teng (ikkilik hisobda)?", "1024", ["1000", "512", "2048"]),
    ("Dasturlash tilini koʻrsating:", "Python", ["Photoshop", "Excel", "Chrome"]),
    ("Internetda manzillarni qidirish uchun ishlatiladigan dastur nima deyiladi?", "Brauzer", ["Printer", "Skaner", "Arxivator"]),
    ("Matematikada 0 ga boʻlish mumkinmi?", "Yoʻq, mumkin emas", ["Ha, natija 0", "Ha, natija 1", "Faqat musbat sonlar uchun"]),
    ("Ikki sonning yigʻindisi 15, ayirmasi 3 boʻlsa, katta son nechaga teng?", "9", ["6", "12", "10"]),
    ("Aylana uzunligi formulasi qaysi?", "2πr", ["πr²", "πd²", "4πr"]),
    ("Gʻalaba bayrami Oʻzbekistonda qaysi sanada nishonlanadi?", "9-may", ["1-may", "8-mart", "21-mart"]),
    ("Navroz bayrami qaysi sanada nishonlanadi?", "21-mart", ["1-yanvar", "9-may", "1-sentabr"]),
    ("Fizikada kuch qaysi birlikda oʻlchanadi?", "Nyuton", ["Joul", "Paskal", "Kelvin"]),
    ("Qaysi element suvning tarkibida mavjud?", "Vodorod", ["Geliy", "Uglerod", "Temir"]),
]
assert len(BANK) >= 50, len(BANK)


def _rows() -> list[tuple[str, list[str], int]]:
    """Har bir savol uchun (matn, 4 ta variant, to'g'ri indeks); to'g'ri javob o'rni aylanib turadi."""
    out = []
    for i, (q, right, wrong) in enumerate(BANK):
        pos = i % 4
        opts = list(wrong)
        opts.insert(pos, right)
        out.append((q, opts, pos))
    return out


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Namuna test bazasi (to'g'ri javob yonida + belgisi)")
    for n, (q, opts, pos) in enumerate(_rows(), start=1):
        doc.add_paragraph(f"{n}. {q}")
        for i, o in enumerate(opts):
            doc.add_paragraph(f"{'ABCD'[i]}) {o}{' +' if i == pos else ''}")
        doc.add_paragraph("")
    doc.save(path)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Savollar"
    ws.append(["Savol", "A", "B", "C", "D"])
    for q, opts, pos in _rows():
        ws.append([q] + [o + (" +" if i == pos else "") for i, o in enumerate(opts)])
    ws.column_dimensions["A"].width = 70
    wb.save(path)


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    make_docx(here / "sample.docx")
    make_xlsx(here / "sample.xlsx")
    print(f"Yaratildi: {here / 'sample.docx'} va {here / 'sample.xlsx'} ({len(BANK)} savol)")
