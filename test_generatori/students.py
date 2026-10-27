"""O'quvchilar ro'yxatini o'qish: F.I.Sh. | Sinf | Maktab (.xlsx, .csv, .txt, .docx)."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

from generator import MAX_STUDENTS
from models import Student
from parser import ParseError, _cell_text

_NAME_KEYS = ("fish", "fio", "ism", "ismi", "familiya", "familiyaism", "familiyavaism", "familiyaismi", "oquvchi",
              "oquvchinging", "oquvchiningfish", "oquvchiningismi", "oquvchining", "talaba", "name", "fullname",
              "student", "фио", "имя", "фамилияимя", "ученик", "ф.и.о", "исм", "ўқувчи", "талаба")
_CLASS_KEYS = ("sinf", "sinfi", "class", "класс", "синф", "guruh")
_SCHOOL_KEYS = ("maktab", "maktabi", "school", "школа", "мактаб")


def _norm(s: str) -> str:
    t = s.casefold()
    for ch in ("ʻ", "ʼ", "‘", "’", "`", "'"):
        t = t.replace(ch, "")
    return re.sub(r"[\s\.\-_:№#()]+", "", t)


def _rows_to_students(rows: list[list[str]]) -> list[Student]:
    rows = [[(c or "").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    cols = {"name": 0, "klass": 1, "school": 2}
    head = [_norm(c) for c in rows[0]]
    found: dict[str, int] = {}
    for i, h in enumerate(head):
        if not h:
            continue
        for key, words in (("name", _NAME_KEYS), ("klass", _CLASS_KEYS), ("school", _SCHOOL_KEYS)):
            if key not in found and h in {_norm(w) for w in words}:
                found[key] = i
                break
    if found:                       # birinchi qator sarlavha
        cols = {k: found.get(k, -1) for k in cols}
        if "name" not in found:
            cols["name"] = next((i for i in range(len(head)) if i not in found.values()), 0)
        rows = rows[1:]
    out = []
    for r in rows:
        def get(k: str) -> str:
            i = cols[k]
            return r[i] if 0 <= i < len(r) else ""
        name = get("name")
        if not name:
            continue
        out.append(Student(name=" ".join(name.split()), klass=" ".join(get("klass").split()),
                           school=" ".join(get("school").split())))
    return out


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_students(path: str | Path) -> list[Student]:
    """Ro'yxat faylidan o'quvchilarni o'qiydi. Birinchi ustun F.I.Sh., so'ng Sinf, Maktab."""
    p = Path(path)
    if not p.is_file():
        raise ParseError(f"O'quvchilar ro'yxati fayli topilmadi: {p}")
    ext = p.suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            try:
                rows = [[_cell_text(v) for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
            finally:
                wb.close()
        elif ext in (".csv", ".txt"):
            text = _read_text(p)
            sample = text[:2000]
            delim = max(";,\t|", key=sample.count) if any(d in sample for d in ";,\t|") else ","
            rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        elif ext == ".docx":
            from docx import Document
            doc = Document(str(p))
            rows = []
            for t in doc.tables:
                for r in t.rows:
                    seen, cells = set(), []
                    for c in r.cells:
                        if id(c._tc) not in seen:
                            seen.add(id(c._tc))
                            cells.append(c.text)
                    rows.append(cells)
            if not rows:
                for para in doc.paragraphs:
                    parts = re.split(r"\s*[;\t|]\s*", para.text.strip())
                    rows.append(parts)
        else:
            raise ParseError("O'quvchilar ro'yxati .xlsx, .csv, .txt yoki .docx bo'lishi kerak.")
    except ParseError:
        raise
    except PermissionError as e:
        raise ParseError("Ro'yxat faylini o'qib bo'lmadi. Faylni yoping va qayta urinib ko'ring.") from e
    except (zipfile.BadZipFile, KeyError) as e:
        raise ParseError("Ro'yxat fayli buzilgan yoki noto'g'ri formatda.") from e
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"Ro'yxat faylini o'qib bo'lmadi: {e}") from e
    students = _rows_to_students(rows)
    if not students:
        raise ParseError("O'quvchilar ro'yxatida birorta ham F.I.Sh. topilmadi.")
    if len(students) > MAX_STUDENTS:
        raise ParseError(f"Ro'yxatda {len(students)} ta o'quvchi bor; eng ko'pi bilan {MAX_STUDENTS} ta mumkin.")
    return students
