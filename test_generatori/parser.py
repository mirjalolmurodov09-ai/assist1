"""Word (.docx) va Excel (.xlsx) fayllaridan test savollarini o'qish.

To'g'ri javob variant matnining boshida yoki oxirida "+" belgisi bilan ko'rsatiladi.
Variant belgilari: A) A. A: (A) a) — lotin va kirill harflarida.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Iterable

from models import LETTERS, Issue, ParseResult, Question


class ParseError(Exception):
    """Faylni umuman o'qib bo'lmaydigan holat (foydalanuvchiga ko'rsatiladi)."""


# --------------------------------------------------------------------------
# Variant belgilari
# --------------------------------------------------------------------------
# belgi -> u egallashi mumkin bo'lgan o'rinlar (0=A, 1=B, 2=C, 3=D, 4=E, 5=F).
# Kirill "В" ikki xil o'qiladi: lotincha B (1) yoki "А Б В Г" tizimidagi uchinchi (2).
_LABEL_POS: dict[str, frozenset[int]] = {}


def _reg(chars: str, positions: Iterable[int]) -> None:
    for ch in chars:
        _LABEL_POS[ch] = frozenset(positions)


_reg("Aa", [0]); _reg("Bb", [1]); _reg("Cc", [2]); _reg("Dd", [3]); _reg("Ee", [4]); _reg("Ff", [5])
_reg("Аа", [0])            # kirill А
_reg("Бб", [1])            # kirill Б
_reg("Вв", [1, 2])         # kirill В
_reg("Сс", [2])            # kirill С (lotin C ga o'xshash)
_reg("Гг", [3])            # kirill Г
_reg("Дд", [3, 4])         # kirill Д
_reg("Ее", [4])            # kirill Е

_CLASS = "[" + "".join(re.escape(c) for c in _LABEL_POS) + "]"
# belgi: boshida yoki bo'shliqdan keyin; ixtiyoriy "(", harf, ")" "." yoki ":"
_MARK_RE = re.compile(rf"(?<!\S)\(?({_CLASS})\s?([\)\.:])")
_LEAD_PLUS_MARK_RE = re.compile(rf"^\+\s*(?=\(?{_CLASS}\s?[\)\.:])")
_QNUM_RE = re.compile(r"^\s*(\d{1,4})\s*(?:\)|\.(?!\d)|[-–—:])\s*(.*)$")
_TRAIL_PLUS_RE = re.compile(r"(?<!\+)\+\s*$")
_LEAD_PLUS_RE = re.compile(r"^\+(?=\s|[^\W\d_])")
_LINE_SPLIT_RE = re.compile(r"[\r\n\x0b\x0c  ]+")


def _clean(s: str) -> str:
    """Ko'rinmas belgilarni olib tashlaydi, bo'shliqlarni tartibga soladi."""
    s = s.replace(" ", " ").replace("\t", " ")
    for ch in ("​", "‌", "‍", "﻿", "­"):
        s = s.replace(ch, "")
    return re.sub(r"[ ]{2,}", " ", s).strip()


def _split_markers(line: str, expected: int) -> list[tuple[int | None, str]]:
    """Qatorni variant belgilari bo'yicha bo'ladi.

    `expected` — navbatdagi kutilayotgan variant o'rni (0=A). Qaytadi:
    [(None, belgidan oldingi matn)?, (o'rin, matn), ...].
    Qator o'rtasidagi belgi faqat ")" bilan va kamida ikkita ketma-ket variant bo'lsa qabul qilinadi.
    """
    accepted: list[tuple[int, int, int]] = []
    e = expected
    for m in _MARK_RE.finditer(line):
        if e not in _LABEL_POS[m.group(1)]:
            continue
        if m.start() != 0 and m.group(2) != ")":
            continue
        accepted.append((m.start(), m.end(), e))
        e += 1
    if accepted and accepted[0][0] != 0 and len(accepted) < 2:
        accepted = []
    if not accepted:
        return [(None, line)]
    out: list[tuple[int | None, str]] = []
    if accepted[0][0] > 0:
        out.append((None, line[: accepted[0][0]].strip()))
    for i, (_, end, idx) in enumerate(accepted):
        nxt = accepted[i + 1][0] if i + 1 < len(accepted) else len(line)
        out.append((idx, line[end:nxt].strip()))
    return out


def _strip_plus(text: str) -> tuple[str, bool]:
    """Variant matnidan "+" belgisini ajratadi. (tozalangan matn, to'g'ri belgisi bormi)."""
    t = text.strip()
    plus = False
    if _LEAD_PLUS_RE.match(t):
        plus = True
        t = t[1:].strip()
    m = _TRAIL_PLUS_RE.search(t)
    if m:
        plus = True
        t = t[: m.start()].strip()
    return t, plus


# --------------------------------------------------------------------------
# Savolni tekshirish va yig'ish
# --------------------------------------------------------------------------
def _excerpt(text: str, n: int = 45) -> str:
    t = " ".join(text.split())
    return t if len(t) <= n else t[: n - 1] + "…"


def _build_question(
    where: str, text: str, options: list[str], plus: list[bool],
    source_no: int, issues: list[Issue], forced_correct: int | None = None,
) -> Question | None:
    """Savolni tekshiradi; xato bo'lsa `issues` ga yozadi va None qaytaradi."""
    text = text.strip()
    bad = False

    def err(msg: str) -> None:
        nonlocal bad
        bad = True
        issues.append(Issue("xato", where, msg))

    if not text:
        err("savol matni bo'sh")
    if len(options) != 4:
        err(f"variantlar soni {len(options)} ta (aynan 4 ta: A, B, C, D bo'lishi kerak)")
    for i, o in enumerate(options):
        if not o.strip():
            err(f"{LETTERS[i] if i < 4 else i + 1}) variant matni bo'sh")
    marked = [i for i, p in enumerate(plus) if p]
    correct: int | None = None
    if forced_correct is not None:
        if marked and marked != [forced_correct]:
            err("«To'g'ri javob» ustuni va «+» belgisi bir-biriga zid")
        correct = forced_correct
    elif len(marked) == 0:
        err("to'g'ri javob («+» belgisi) ko'rsatilmagan")
    elif len(marked) > 1:
        err("bir nechta variantda «+» belgisi bor: " + ", ".join(LETTERS[i] for i in marked if i < 4))
    else:
        correct = marked[0]
    if bad or correct is None or correct >= 4:
        return None
    norm = [" ".join(o.casefold().split()) for o in options]
    if len(set(norm)) < 4:
        issues.append(Issue("ogohlantirish", where, "ba'zi variantlar bir xil matnga ega"))
    return Question(text=text, options=tuple(options), correct=correct, source_no=source_no)  # type: ignore[arg-type]


def _dedupe(questions: list[Question], issues: list[Issue]) -> list[Question]:
    """Aynan takrorlangan savollarni olib tashlaydi (ogohlantirish bilan)."""
    seen: dict[tuple, int] = {}
    out: list[Question] = []
    for q in questions:
        key = (
            " ".join(q.text.casefold().split()),
            frozenset(" ".join(o.casefold().split()) for o in q.options),
            " ".join(q.options[q.correct].casefold().split()),
        )
        if key in seen:
            issues.append(Issue(
                "ogohlantirish", f"{q.source_no}-savol",
                f"{seen[key]}-savolning aynan takrori — o'tkazib yuborildi («{_excerpt(q.text)}»)"))
            continue
        seen[key] = q.source_no
        out.append(q)
    return out


# --------------------------------------------------------------------------
# Matn (Word) qatorlarini o'qish
# --------------------------------------------------------------------------
class _Block:
    """Hozir yig'ilayotgan savol."""

    def __init__(self, number: int | None) -> None:
        self.number = number
        self.q_lines: list[str] = []
        self.opts: list[dict] = []  # {"parts": [str], "plus": bool}


class _TextParser:
    """Qatorma-qator holat mashinasi: savol -> variantlar -> keyingi savol."""

    def __init__(self) -> None:
        self.cur: _Block | None = None
        self.blocks: list[tuple[_Block, int]] = []  # (blok, tartib raqami)
        self.questions: list[Question] = []
        self.issues: list[Issue] = []
        self.seq = 0
        self.found = 0

    # ---- yordamchilar
    def _target_append(self, text: str) -> None:
        cur = self.cur
        assert cur is not None
        if not text:
            return
        if cur.opts:
            cur.opts[-1]["parts"].append(text)
        else:
            cur.q_lines.append(text)

    def _feed(self, content: str, lead_plus: bool = False) -> None:
        """`content` ni joriy blokka qo'shadi (variant belgilari bo'yicha bo'lib)."""
        cur = self.cur
        assert cur is not None
        for idx, txt in _split_markers(content, len(cur.opts)):
            if idx is None:
                self._target_append(txt)
            else:
                cur.opts.append({"parts": [txt] if txt else [], "plus": lead_plus})
                lead_plus = False

    def flush(self) -> None:
        """Joriy blokni yakunlaydi."""
        cur = self.cur
        self.cur = None
        if cur is None:
            return
        text = "\n".join(cur.q_lines).strip()
        if cur.number is None and not cur.opts:
            if text:
                self.issues.append(Issue(
                    "ogohlantirish", "matn",
                    f"variantsiz va raqamsiz matn savol deb olinmadi, o'tkazib yuborildi: «{_excerpt(text)}»"))
            return
        self.seq += 1
        self.found += 1
        where = (f"{cur.number}-savol" if cur.number is not None else f"{self.seq}-blok")
        where += f" («{_excerpt(text)}»)" if text else ""
        options: list[str] = []
        plus: list[bool] = []
        for o in cur.opts:
            t, p = _strip_plus(" ".join(o["parts"]))
            options.append(t)
            plus.append(p or o["plus"])
        q = _build_question(where, text, options, plus, self.seq, self.issues)
        if q is not None:
            self.questions.append(q)

    # ---- asosiy
    def feed_line(self, raw: str) -> None:
        line = _clean(raw)
        if not line:
            return
        cur = self.cur
        expect = len(cur.opts) if cur else 0

        # 1) joriy savolning navbatdagi varianti (qator boshida belgi bor)
        if cur is not None:
            lead_plus = False
            probe = line
            if _LEAD_PLUS_MARK_RE.match(probe):
                lead_plus = True
                probe = _LEAD_PLUS_MARK_RE.sub("", probe, count=1)
            parts = _split_markers(probe, expect)
            if parts and parts[0][0] is not None:
                self._feed(probe, lead_plus)
                return

        # 2) raqamli savol boshlanishi
        m = _QNUM_RE.match(line)
        if m:
            num = int(m.group(1))
            new_q = (
                cur is None
                or len(cur.opts) >= 4
                or cur.number is None
                or num == cur.number + 1
            )
            if new_q:
                self.flush()
                self.cur = _Block(num)
                self._feed(m.group(2).strip())
                return

        # 3) raqamsiz yangi savol (4 ta variant to'lgandan keyin kelgan oddiy matn)
        if cur is not None and len(cur.opts) >= 4:
            self.flush()
            self.cur = _Block(None)
            self._feed(line)
            return
        if cur is None:
            self.cur = _Block(None)
            self._feed(line)
            return

        # 4) joriy savol yoki varianti davomi
        self._feed(line)


def _parse_lines(lines: Iterable[str], tp: _TextParser | None = None) -> _TextParser:
    tp = tp or _TextParser()
    for ln in lines:
        tp.feed_line(ln)
    return tp


# --------------------------------------------------------------------------
# Jadval ko'rinishidagi ma'lumot (Excel varag'i yoki Word jadvali)
# --------------------------------------------------------------------------
_Q_HEADERS = {"savol", "savollar", "savolmatni", "question", "questions", "вопрос", "савол", "саволлар", "test", "testlar"}
_ANS_HINTS = ("javob", "togri", "to'g'ri", "to‘g‘ri", "toʻgʻri", "answer", "kalit", "жавоб", "тўғри", "правиль")


def _norm_head(s: str) -> str:
    return re.sub(r"[\s\.\)\(:№#_\-]+", "", s.casefold())


def _label_to_idx(s: str) -> int | None:
    """"A", "b)", "В", "2" kabi qiymatni 0..3 indeksga aylantiradi."""
    t = s.strip().strip(".):(").strip()
    if len(t) != 1:
        return None
    if t in "1234":
        return int(t) - 1
    pos = _LABEL_POS.get(t)
    if not pos:
        return None
    idx = min(pos)
    return idx if idx < 4 else None


def _strip_label(text: str, idx: int) -> str:
    """Yacheyka boshidagi "A)" kabi belgini (o'z o'rniga mos bo'lsa) olib tashlaydi."""
    m = re.match(rf"^\s*\+?\s*\(?({_CLASS})\s?[\)\.:]\s*", text)
    if m and idx in _LABEL_POS[m.group(1)]:
        lead = "+ " if text.lstrip().startswith("+") else ""
        return lead + text[m.end():]
    return text


def parse_grid(rows: list[tuple[int, list[str]]], prefix: str, tp: _TextParser | None = None) -> _TextParser:
    """Jadval qatorlarini savollarga aylantiradi. rows = [(qator raqami, [yacheykalar])]."""
    tp = tp or _TextParser()
    tp.flush()
    rows = [(n, [_clean(c) for c in cells]) for n, cells in rows]
    rows = [(n, c) for n, c in rows if any(c)]
    if not rows:
        return tp

    q_col, opt_cols, ans_col = 0, [1, 2, 3, 4], None
    start = 0
    head = [_norm_head(c) for c in rows[0][1]]
    letter_cols = {}
    for i, h in enumerate(head):
        if h in ("a", "b", "c", "d", "а", "б", "в", "с", "г", "д"):
            idx = {"a": 0, "b": 1, "c": 2, "d": 3, "а": 0, "б": 1, "в": 1, "с": 2, "г": 3, "д": 3}[h]
            letter_cols.setdefault(idx, i)
    is_header = False
    if len(letter_cols) == 4:
        is_header = True
        opt_cols = [letter_cols[i] for i in range(4)]
        first_opt = min(opt_cols)
        q_col = next((i for i, h in enumerate(head) if h in _Q_HEADERS), max(first_opt - 1, 0))
        if q_col in opt_cols:
            q_col = 0
    elif head and head[0] in _Q_HEADERS or (len(head) > 1 and head[1] in _Q_HEADERS):
        is_header = True
        q_col = 0 if head[0] in _Q_HEADERS else 1
        opt_cols = [q_col + 1 + i for i in range(4)]
    if is_header:
        start = 1
        for i, h in enumerate(rows[0][1]):
            low = h.casefold()
            if i != q_col and i not in opt_cols and any(k in low for k in _ANS_HINTS):
                ans_col = i
                break
    else:
        first = rows[0][1]
        off = 1 if (len(first) >= 6 and first[0].strip().isdigit()) else 0
        q_col = off
        opt_cols = [off + 1 + i for i in range(4)]
        if len(first) > off + 5:
            ans_col = off + 5

    for n, cells in rows[start:]:
        def cell(i: int) -> str:
            return cells[i] if 0 <= i < len(cells) else ""

        tp.seq += 1
        tp.found += 1
        text = cell(q_col)
        opts: list[str] = []
        plus: list[bool] = []
        for k, ci in enumerate(opt_cols):
            t, p = _strip_plus(_strip_label(cell(ci), k))
            opts.append(t)
            plus.append(p)
        forced = None
        where = f"{prefix} {n}-qator"
        if ans_col is not None and cell(ans_col):
            forced = _label_to_idx(cell(ans_col))
            if forced is None:
                tp.issues.append(Issue("xato", where, f"«To'g'ri javob» ustunidagi qiymat tushunarsiz: «{cell(ans_col)}»"))
                continue
        q = _build_question(where, text, opts, plus, tp.seq, tp.issues, forced)
        if q is not None:
            tp.questions.append(q)
    return tp


# --------------------------------------------------------------------------
# Fayl turlari
# --------------------------------------------------------------------------
def _cell_text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(path: Path) -> _TextParser:
    """Excel faylining birinchi varag'ini o'qiydi."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError) as e:
        raise ParseError("Fayl buzilgan yoki haqiqiy .xlsx fayl emas.") from e
    try:
        ws = wb.worksheets[0]
        rows = []
        for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_cell_text(v) for v in row]
            if any(cells):
                rows.append((n, cells))
    finally:
        wb.close()
    return parse_grid(rows, "Excel")


def _iter_docx_blocks(doc) -> Iterable[object]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def _table_rows(table) -> list[list[str]]:
    """Jadval qatorlari (birlashtirilgan yacheykalar takrorlanmaydi)."""
    out = []
    for row in table.rows:
        seen: set[int] = set()
        cells = []
        for c in row.cells:
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            cells.append("\n".join(p.text for p in c.paragraphs))
        out.append(cells)
    return out


def parse_docx(path: Path) -> _TextParser:
    """Word hujjatidagi savollarni o'qiydi (paragraflar va jadvallar tartib bilan)."""
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table

    try:
        doc = Document(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError) as e:
        raise ParseError("Fayl buzilgan yoki haqiqiy .docx fayl emas.") from e

    tp = _TextParser()
    for block in _iter_docx_blocks(doc):
        if isinstance(block, Table):
            rows = _table_rows(block)
            wide = sum(1 for r in rows if len([c for c in r if c.strip()]) >= 5)
            if rows and wide >= max(1, len(rows) // 2):
                parse_grid([(i + 1, r) for i, r in enumerate(rows)], "Jadval", tp)
            else:
                for r in rows:
                    for cell in r:
                        for ln in _LINE_SPLIT_RE.split(cell):
                            tp.feed_line(ln)
        else:
            for ln in _LINE_SPLIT_RE.split(block.text):  # type: ignore[attr-defined]
                tp.feed_line(ln)
    tp.flush()
    return tp


def parse_text(text: str) -> ParseResult:
    """Oddiy matndan o'qish (testlar va tekshirish uchun qulay)."""
    tp = _parse_lines(_LINE_SPLIT_RE.split(text))
    tp.flush()
    return _finish(tp)


def _finish(tp: _TextParser) -> ParseResult:
    issues = list(tp.issues)
    questions = _dedupe(tp.questions, issues)
    return ParseResult(questions=questions, issues=issues, total_found=tp.found)


def parse_file(path: str | Path) -> ParseResult:
    """Fayl turiga qarab savollarni o'qiydi. Tushunarli ParseError qaytaradi."""
    p = Path(path)
    if p.name.startswith("~$"):
        raise ParseError("Bu Word/Excel'ning vaqtinchalik (~$) fayli. Asl faylni tanlang.")
    if not p.is_file():
        raise ParseError(f"Fayl topilmadi: {p}")
    ext = p.suffix.lower()
    try:
        if ext == ".docx":
            tp = parse_docx(p)
        elif ext in (".xlsx", ".xlsm"):
            tp = parse_xlsx(p)
        elif ext in (".doc", ".xls"):
            raise ParseError(
                f"{ext} eski format qo'llab-quvvatlanmaydi. Faylni Word/Excelda "
                "«Boshqacha saqlash» orqali .docx yoki .xlsx qilib saqlang.")
        else:
            raise ParseError("Faqat .docx va .xlsx fayllar qo'llab-quvvatlanadi.")
    except ParseError:
        raise
    except PermissionError as e:
        raise ParseError("Faylni o'qib bo'lmadi (ruxsat yo'q yoki fayl band). Faylni yoping va qayta urinib ko'ring.") from e
    except OSError as e:
        raise ParseError(f"Faylni o'qishda xatolik: {e}") from e
    except Exception as e:  # buzilgan fayllar uchun
        raise ParseError(f"Faylni o'qib bo'lmadi: {e}") from e

    res = _finish(tp)
    if res.total_found == 0:
        raise ParseError("Faylda birorta ham savol topilmadi. Format namunasini README.md dan ko'ring.")
    return res
