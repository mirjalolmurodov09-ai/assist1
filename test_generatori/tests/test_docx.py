import re
import shutil
import subprocess

import pytest
from docx import Document
from docx.enum.section import WD_SECTION

from docx_builder import build_docx
from generator import generate_tests
from helpers import make_bank
from models import Question
from parser import parse_file
from pdf_builder import PdfMeta


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_every_question_and_option_present_no_plus(tmp_path):
    res = generate_tests(make_bank(60), 5, 20, seed=1)
    out = tmp_path / "t.docx"
    build_docx(res.tests, out, PdfMeta("Algebra testi", "9-sinf"))
    assert out.exists() and not (tmp_path / "t.docx.part").exists()
    doc = Document(str(out))
    txt = re.sub(r"[ \t]+", " ", all_text(doc))
    assert "+" not in txt
    assert "Algebra testi" in txt and "9-sinf" in txt
    for t in res.tests:
        assert t.test_id in txt
        for n, it in enumerate(t.items, start=1):
            assert it.text in txt
            for o in it.options:
                assert o in txt


def test_one_section_per_test_with_ids_in_header_and_footer(tmp_path):
    res = generate_tests(make_bank(30), 4, 10, seed=2)
    out = tmp_path / "s.docx"
    build_docx(res.tests, out, PdfMeta("X"))
    doc = Document(str(out))
    assert len(doc.sections) == 4
    for sec, t in zip(doc.sections, res.tests):
        assert sec.different_first_page_header_footer
        assert t.test_id in sec.footer.paragraphs[0].text
        assert t.test_id in sec.header.paragraphs[0].text          # keyingi betlarda mayda ID
        assert not sec.footer.is_linked_to_previous and not sec.header.is_linked_to_previous or sec is doc.sections[0]
        assert (sec.page_width.mm, sec.page_height.mm) == pytest.approx((210, 297), abs=0.5)


def test_even_pages_uses_odd_page_section_breaks(tmp_path):
    res = generate_tests(make_bank(30), 3, 10, seed=3)
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    build_docx(res.tests, a, even_pages=True)
    build_docx(res.tests, b, even_pages=False)
    starts_a = [s.start_type for s in Document(str(a)).sections]
    starts_b = [s.start_type for s in Document(str(b)).sections]
    assert starts_a[1:] == [WD_SECTION.ODD_PAGE] * 2
    assert starts_b[1:] == [WD_SECTION.NEW_PAGE] * 2
    # oddiy rejimda har bir testda bet raqami 1 dan boshlanadi
    doc_b, doc_a = Document(str(b)), Document(str(a))
    assert doc_b.element.xml.count("w:pgNumType") == 3
    assert all("PAGE" in s.footer._element.xml for s in doc_b.sections)
    # toq-bet rejimida raqam qayta boshlanmaydi va kolontitulda faqat ID bor
    assert "w:pgNumType" not in doc_a.element.xml
    assert all("PAGE" not in s.footer._element.xml for s in doc_a.sections)


def test_sectPr_children_in_schema_order(tmp_path):
    """Word «o'qib bo'lmaydigan kontent» deb rad etmasligi uchun sectPr tartibi to'g'ri bo'lishi shart."""
    res = generate_tests(make_bank(20), 2, 5, seed=1)
    out = tmp_path / "o.docx"
    build_docx(res.tests, out, even_pages=False)
    rank = {"headerReference": 0, "footerReference": 0, "type": 1, "pgSz": 2, "pgMar": 3, "pgNumType": 4,
            "cols": 5, "titlePg": 6, "docGrid": 7}           # header/footer havolalari o'zaro erkin tartibda
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr"
    for sect in Document(str(out)).element.body.iter(ns):
        idx = [rank[c.tag.rsplit("}", 1)[-1]] for c in sect if c.tag.rsplit("}", 1)[-1] in rank]
        assert idx == sorted(idx) and 4 in idx


def test_special_characters_and_control_chars(tmp_path):
    q = Question("Ўзбекистон <b>&</b> «Oʻzbek»\x00\x0b ikki\nqator?",
                 ("Тошкент\x07", "A & B <i>", "\"qo'sh\"", "juda " * 60), 0, 1)
    res = generate_tests([q, *make_bank(3)], 2, 4, seed=1, shuffle_options=False)
    out = tmp_path / "sp.docx"
    build_docx(res.tests, out)
    txt = all_text(Document(str(out)))
    assert "Ўзбекистон <b>&</b> «Oʻzbek»" in txt and "Тошкент" in txt and "A & B <i>" in txt
    assert "\x00" not in txt and "\x07" not in txt


def test_our_parser_recognises_structure_of_generated_docx(tmp_path):
    res = generate_tests(make_bank(40), 1, 12, seed=1)
    out = tmp_path / "r.docx"
    build_docx(res.tests, out)
    r = parse_file(out)
    assert r.total_found == 12                 # 12 ta savol blok sifatida taniladi (to'g'ri javob yo'qligi uchun xato)
    assert all("ko'rsatilmagan" in e.message for e in r.errors)


def test_failed_build_removes_partial_file(tmp_path):
    res = generate_tests(make_bank(10), 2, 3, seed=1)
    out = tmp_path / "f.docx"

    def boom(d, t):
        raise RuntimeError("to'xtatildi")

    with pytest.raises(RuntimeError):
        build_docx(res.tests, out, progress=boom)
    assert not out.exists() and not (tmp_path / "f.docx.part").exists()


def test_locked_file_message(tmp_path, monkeypatch):
    import os
    res = generate_tests(make_bank(10), 1, 3, seed=1)

    def deny(*a, **k):
        raise PermissionError("busy")
    monkeypatch.setattr(os, "replace", deny)
    with pytest.raises(PermissionError, match="yoping"):
        build_docx(res.tests, tmp_path / "l.docx")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice yo'q")
def test_renders_in_libreoffice(tmp_path):
    from pypdf import PdfReader
    res = generate_tests(make_bank(60), 3, 40, seed=1)
    out = tmp_path / "lo.docx"
    build_docx(res.tests, out, PdfMeta("LO"), even_pages=False)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(out)],
                   check=True, capture_output=True, timeout=180)
    pdf = tmp_path / "lo.pdf"
    texts = [p.extract_text() for p in PdfReader(str(pdf)).pages]
    flat = "\n".join(texts)
    assert len(texts) >= 6
    for t in res.tests:
        assert t.test_id in flat


def test_paragraph_and_run_properties_in_schema_order(tmp_path):
    """Word elementlar tartibiga talabchan: pPr va rPr ichidagi tartib sxemaga mos bo'lishi kerak."""
    res = generate_tests(make_bank(30), 3, 10, seed=1)
    out = tmp_path / "ord.docx"
    build_docx(res.tests, out, PdfMeta("T", "S"), even_pages=True)
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    ppr_order = ["pStyle", "keepNext", "keepLines", "pBdr", "tabs", "spacing", "ind", "jc", "rPr", "sectPr"]
    rpr_order = ["rFonts", "b", "bCs", "sz", "szCs"]
    doc = Document(str(out))
    seen_sect_in_ppr = 0
    for ppr in doc.element.iter(w + "pPr"):
        names = [c.tag[len(w):] for c in ppr]
        idx = [ppr_order.index(n) for n in names if n in ppr_order]
        assert idx == sorted(idx), names
        seen_sect_in_ppr += "sectPr" in names
    assert seen_sect_in_ppr == 2                 # 3 ta test = 2 ta paragraf ichidagi bo'lim uzilishi + oxirgi
    for rpr in doc.element.iter(w + "rPr"):
        names = [c.tag[len(w):] for c in rpr]
        idx = [rpr_order.index(n) for n in names if n in rpr_order]
        assert idx == sorted(idx), names
    # shrift hujjat uslubida (tema emas) berilgan
    normal = doc.styles["Normal"].element.xml
    assert 'w:ascii="Arial"' in normal
    assert "asciiTheme" not in doc.styles.element.xml
