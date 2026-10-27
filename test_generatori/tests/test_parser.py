import pytest
from docx import Document
from openpyxl import Workbook

from parser import ParseError, parse_file, parse_text


def q(res, i=0):
    x = res.questions[i]
    return x.text, list(x.options), x.correct


def test_standard_format():
    r = parse_text("""
1. Oʻzbekiston poytaxti qayer?
A) Samarqand
B) Toshkent +
C) Buxoro
D) Xiva

2. 2+2 nechaga teng?
A) 3
B) 4 +
C) 5
D) 6
""")
    assert len(r.questions) == 2 and not r.errors
    assert q(r, 0) == ("Oʻzbekiston poytaxti qayer?", ["Samarqand", "Toshkent", "Buxoro", "Xiva"], 1)
    assert q(r, 1) == ("2+2 nechaga teng?", ["3", "4", "5", "6"], 1)


@pytest.mark.parametrize("fmt", ["{L})", "{L}.", "{L}:", "({L})", "{l})"])
def test_label_styles(fmt):
    lines = ["1. Savol?"]
    for i, L in enumerate("ABCD"):
        lab = fmt.format(L=L, l=L.lower())
        lines.append(f"{lab} javob{i}" + (" +" if i == 2 else ""))
    r = parse_text("\n".join(lines))
    assert not r.errors and q(r) == ("Savol?", [f"javob{i}" for i in range(4)], 2)


def test_cyrillic_labels_latin_scheme():
    r = parse_text("1. Savol?\nА) bir\nВ) ikki\nС) uch +\nД) toʻrt")
    assert not r.errors and q(r) == ("Savol?", ["bir", "ikki", "uch", "toʻrt"], 2)


def test_cyrillic_labels_abvg_scheme():
    r = parse_text("1. Савол?\nА) бир\nБ) икки\nВ) уч +\nГ) тўрт")
    assert not r.errors and q(r) == ("Савол?", ["бир", "икки", "уч", "тўрт"], 2)


def test_plus_at_start_and_end_and_before_label():
    r = parse_text("1. S?\nA) bir\n+B) ikki\nC) uch\nD) to'rt\n"
                   "2. S2?\nA) bir\nB) + ikki\nC) uch\nD) to'rt\n"
                   "3. S3?\nA) bir\nB) ikki\nC) uch\nD) to'rt+")
    assert [x.correct for x in r.questions] == [1, 1, 3]
    assert r.questions[0].options[1] == "ikki" and r.questions[1].options[1] == "ikki"
    assert r.questions[2].options[3] == "to'rt"


def test_plus_never_in_options_text():
    r = parse_text("1. S?\nA) 1\nB) 2 +\nC) 3\nD) 4")
    assert all("+" not in o for o in r.questions[0].options)


def test_cplusplus_not_mistaken():
    r = parse_text("1. Til?\nA) C++\nB) Java +\nC) Go\nD) C#")
    assert not r.errors and q(r)[1][0] == "C++" and q(r)[2] == 1


def test_inline_options():
    r = parse_text("1. Qaysi? A) 3 B) 4+ C) 5 D) 6")
    assert not r.errors and q(r) == ("Qaysi?", ["3", "4", "5", "6"], 1)
    r = parse_text("1. Savol?\nA) 3 B) 4 C) 5+ D) 6")
    assert not r.errors and q(r) == ("Savol?", ["3", "4", "5", "6"], 2)
    r = parse_text("1. Savol?\nA) 3 B) 4\nC) 5 D) 6 +")
    assert not r.errors and q(r)[2] == 3


def test_question_numbering_styles():
    r = parse_text("1) A?\nA) a\nB) b +\nC) c\nD) d\n2- B?\nA) a\nB) b\nC) c +\nD) d\n3.C?\nA) a +\nB) b\nC) c\nD) d")
    assert [x.text for x in r.questions] == ["A?", "B?", "C?"] and not r.errors


def test_multiline_question_and_option():
    r = parse_text("1. Birinchi qator\nikkinchi qator?\nA) uzun\njavob davomi\nB) b +\nC) c\nD) d")
    assert not r.errors
    assert r.questions[0].text == "Birinchi qator\nikkinchi qator?"
    assert r.questions[0].options[0] == "uzun javob davomi"


def test_decimal_continuation_not_new_question():
    r = parse_text("1. Hisoblang:\n3.5 + 1.5 = ?\nA) 4\nB) 5 +\nC) 6\nD) 7")
    assert len(r.questions) == 1 and "3.5" in r.questions[0].text


def test_blank_lines_and_spaces_ignored():
    r = parse_text("\n\n  1.   Savol?   \n\n  A)   a  \n\nB)  b +\n\n\nC) c\nD) d\n\n")
    assert not r.errors and q(r) == ("Savol?", ["a", "b", "c", "d"], 1)


def test_no_numbers_autonumbered_doc():
    r = parse_text("Birinchi savol?\nA) a\nB) b +\nC) c\nD) d\nIkkinchi savol?\nA) a +\nB) b\nC) c\nD) d")
    assert [x.text for x in r.questions] == ["Birinchi savol?", "Ikkinchi savol?"]


def test_header_text_ignored_with_warning():
    r = parse_text("Matematika testi\n1. S?\nA) a +\nB) b\nC) c\nD) d")
    assert len(r.questions) == 1 and r.warnings


def test_errors_reported_with_location():
    r = parse_text("""1. Toʻgʻri belgisiz?
A) a
B) b
C) c
D) d
2. Ikki plus?
A) a +
B) b +
C) c
D) d
3. Uchta variant?
A) a +
B) b
C) c
4. Yaxshi?
A) a +
B) b
C) c
D) d
5. Boʻsh variant?
A) a +
B)
C) c
D) d""")
    assert len(r.questions) == 1 and r.questions[0].text == "Yaxshi?"
    errs = [str(e) for e in r.errors]
    assert len(errs) == 4
    assert any("1-savol" in e and "ko'rsatilmagan" in e for e in errs)
    assert any("2-savol" in e and "bir nechta" in e for e in errs)
    assert any("3-savol" in e and "3 ta" in e for e in errs)
    assert any("5-savol" in e and "bo'sh" in e for e in errs)


def test_duplicates_warned_and_removed():
    block = "{n}. Bir xil?\nA) a +\nB) b\nC) c\nD) d\n"
    r = parse_text(block.format(n=1) + block.format(n=2))
    assert len(r.questions) == 1 and any("takr" in w.message for w in r.warnings)


def test_same_text_different_options_not_duplicate():
    r = parse_text("1. Qaysi?\nA) a +\nB) b\nC) c\nD) d\n2. Qaysi?\nA) x +\nB) y\nC) z\nD) w")
    assert len(r.questions) == 2


# ---------------------------------------------------------------- Word fayllari
def make_docx(path, paragraphs, table_rows=None):
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table_rows:
        t = d.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = t.add_row().cells
            for c, v in zip(cells, row):
                c.text = v
    d.save(path)


def test_docx_with_soft_line_breaks(tmp_path):
    p = tmp_path / "a.docx"
    d = Document()
    para = d.add_paragraph()
    para.add_run("1. Savol?")
    para.add_run().add_break()
    para.add_run("A) a")
    para.add_run().add_break()
    para.add_run("B) b +")
    para.add_run().add_break()
    para.add_run("C) c")
    para.add_run().add_break()
    para.add_run("D) d")
    d.save(p)
    r = parse_file(p)
    assert not r.errors and q(r) == ("Savol?", ["a", "b", "c", "d"], 1)


def test_docx_table_flat_lines(tmp_path):
    p = tmp_path / "t.docx"
    make_docx(p, [], [["1. Jadvaldagi savol?"], ["A) a"], ["B) b +"], ["C) c"], ["D) d"]])
    r = parse_file(p)
    assert not r.errors and q(r) == ("Jadvaldagi savol?", ["a", "b", "c", "d"], 1)


def test_docx_table_grid(tmp_path):
    p = tmp_path / "g.docx"
    make_docx(p, ["Sarlavha"], [["Savol", "A", "B", "C", "D"], ["Birinchi?", "a", "b +", "c", "d"],
                                ["Ikkinchi?", "a +", "b", "c", "d"]])
    r = parse_file(p)
    assert [x.correct for x in r.questions] == [1, 0] and not r.errors


def test_docx_paragraphs_and_table_in_order(tmp_path):
    p = tmp_path / "m.docx"
    make_docx(p, ["1. Oddiy?", "A) a +", "B) b", "C) c", "D) d"],
              [["2. Jadvalda?"], ["A) a"], ["B) b"], ["C) c +"], ["D) d"]])
    r = parse_file(p)
    assert [x.text for x in r.questions] == ["Oddiy?", "Jadvalda?"]


# ---------------------------------------------------------------- Excel fayllari
def make_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_xlsx_with_header_and_blank_rows(tmp_path):
    p = tmp_path / "a.xlsx"
    make_xlsx(p, [["Savol", "A", "B", "C", "D"], ["S1?", "a", "b +", "c", "d"], [None] * 5,
                  ["S2?", "a +", "b", "c", "d"], [None, None]])
    r = parse_file(p)
    assert [x.text for x in r.questions] == ["S1?", "S2?"] and not r.errors


def test_xlsx_without_header_numbers_and_plus_variants(tmp_path):
    p = tmp_path / "b.xlsx"
    make_xlsx(p, [["S1?", 1, 2, "3 +", 4], ["S2?", "+ a", "b", "c", "d"]])
    r = parse_file(p)
    assert not r.errors and q(r, 0) == ("S1?", ["1", "2", "3", "4"], 2) and r.questions[1].correct == 0


def test_xlsx_answer_column(tmp_path):
    p = tmp_path / "c.xlsx"
    make_xlsx(p, [["Savol", "A", "B", "C", "D", "To'g'ri javob"], ["S1?", "a", "b", "c", "d", "C"],
                  ["S2?", "a", "b", "c", "d", "b)"], ["S3?", "a", "b", "c", "d", "Z"]])
    r = parse_file(p)
    assert [x.correct for x in r.questions] == [2, 1] and len(r.errors) == 1


def test_xlsx_errors_have_row_numbers(tmp_path):
    p = tmp_path / "d.xlsx"
    make_xlsx(p, [["Savol", "A", "B", "C", "D"], ["S1?", "a", "b", "c", "d"], ["S2?", "a +", "b", "c", None]])
    r = parse_file(p)
    errs = [str(e) for e in r.errors]
    assert any("2-qator" in e for e in errs) and any("3-qator" in e for e in errs) and not r.questions


def test_xlsx_option_labels_stripped(tmp_path):
    p = tmp_path / "e.xlsx"
    make_xlsx(p, [["S1?", "A) a", "B) b +", "C) c", "D) d"]])
    assert q(parse_file(p)) == ("S1?", ["a", "b", "c", "d"], 1)


# ---------------------------------------------------------------- Fayl xatolari
def test_missing_and_bad_files(tmp_path):
    with pytest.raises(ParseError, match="topilmadi"):
        parse_file(tmp_path / "yoq.docx")
    bad = tmp_path / "bad.docx"
    bad.write_text("bu docx emas")
    with pytest.raises(ParseError):
        parse_file(bad)
    old = tmp_path / "old.doc"
    old.write_bytes(b"x")
    with pytest.raises(ParseError, match="eski format"):
        parse_file(old)
    txt = tmp_path / "a.txt"
    txt.write_text("x")
    with pytest.raises(ParseError):
        parse_file(txt)
    empty = tmp_path / "empty.docx"
    Document().save(empty)
    with pytest.raises(ParseError, match="topilmadi"):
        parse_file(empty)


def test_cyrillic_and_spaces_in_path(tmp_path):
    d = tmp_path / "Мактаб папка"
    d.mkdir()
    p = d / "тест файл.docx"
    make_docx(p, ["1. S?", "A) a +", "B) b", "C) c", "D) d"])
    assert len(parse_file(p).questions) == 1
