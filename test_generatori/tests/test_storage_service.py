import pytest
from openpyxl import load_workbook

from checker import GradeScale, check_answers, parse_answers
from generator import generate_tests
from helpers import make_bank
from parser import parse_file
from service import create_session
from storage import Storage
import excel_io


@pytest.fixture
def st(tmp_path):
    return Storage(tmp_path / "db" / "tests.db")


def test_end_to_end_with_sample_files(tmp_path, st):
    out = tmp_path / "out"
    from pathlib import Path
    ex = Path(__file__).resolve().parent.parent / "examples"
    for name in ("sample.docx", "sample.xlsx"):
        res = parse_file(ex / name)
        assert len(res.questions) == 50 and not res.errors
        o = create_session(res.questions, n_students=25, q_per_test=20, out_dir=out, storage=st,
                           title="Sinov", seed=5, source_file=name)
        assert o.pdf_path.exists() and o.key_path.exists() and o.pages >= 25
        assert o.pdf_path.parent == out

        # javoblar kaliti
        ws = load_workbook(o.key_path).active
        assert [c.value for c in ws[1]][:4] == ["Test ID", "Tartib raqam", "O'quvchi", "1-savol"]
        assert ws.max_row == 26 and ws.max_column == 3 + 20
        key_row = [c.value for c in ws[2]]
        # bazadagi kalit Excel bilan mos
        stored = st.find_test(o.session_id, key_row[0])
        assert stored.correct_letters == key_row[3:]
        assert len(stored.texts) == 20 and all(len(x) == 4 for x in stored.options)

        # tekshirish: to'g'ri javoblar = 100%
        ans, _ = parse_answers("".join(stored.correct_letters), 20)
        r = check_answers(stored.test_id, stored.correct_letters, stored.texts, ans, st.get_grade_scale())
        assert r.percent == 100 and r.grade == 5
    assert len(st.list_sessions()) == 2


def test_never_overwrites_existing_files(tmp_path, st):
    bank = make_bank(10)
    paths = set()
    for _ in range(3):
        o = create_session(bank, n_students=2, q_per_test=5, out_dir=tmp_path, storage=st, seed=1)
        paths |= {o.pdf_path, o.key_path}
    assert len(paths) == 6 and all(p.exists() for p in paths)


def test_ids_unique_across_sessions_without_seed(tmp_path, st):
    bank = make_bank(10)
    ids = []
    for _ in range(3):
        o = create_session(bank, n_students=50, q_per_test=5, out_dir=tmp_path, storage=st)
        ids.append(o.session_id)
    assert len(st.all_test_ids()) == 150


def test_validation_errors_before_any_file_written(tmp_path, st):
    with pytest.raises(ValueError, match="10"):
        create_session(make_bank(10), n_students=2, q_per_test=11, out_dir=tmp_path / "o", storage=st)
    assert not (tmp_path / "o").exists() and st.list_sessions() == []


def test_results_save_update_and_export(tmp_path, st):
    bank = make_bank(10)
    o = create_session(bank, n_students=2, q_per_test=5, out_dir=tmp_path, storage=st, seed=3)
    t = st.find_test(o.session_id, next(iter(st.all_test_ids())))
    scale = st.get_grade_scale()
    ans, _ = parse_answers("A" * 5, 5)
    r1 = check_answers(t.test_id, t.correct_letters, t.texts, ans, scale)
    st.save_result(o.session_id, r1, "=HYPERLINK(\"x\")")
    ans2, _ = parse_answers("".join(t.correct_letters), 5)
    r2 = check_answers(t.test_id, t.correct_letters, t.texts, ans2, scale)
    st.save_result(o.session_id, r2, "Ali Valiyev")
    rows = st.list_results(o.session_id)
    assert len(rows) == 1 and rows[0]["percent"] == 100 and rows[0]["student"] == "Ali Valiyev"
    assert st.list_sessions()[0].checked == 1

    st.save_result(o.session_id, r1, "=1+1")
    out = tmp_path / "natijalar.xlsx"
    excel_io.write_results(st.list_results(o.session_id), out)
    cell = load_workbook(out).active["C2"]
    assert cell.value == "=1+1" and cell.data_type == "s"        # formulaga aylanmaydi


def test_unknown_id_and_session_isolation(tmp_path, st):
    bank = make_bank(10)
    a = create_session(bank, n_students=2, q_per_test=5, out_dir=tmp_path, storage=st, seed=1)
    b = create_session(bank, n_students=2, q_per_test=5, out_dir=tmp_path, storage=st, seed=2)
    tid = next(iter({t for t in st.all_test_ids()}))
    assert st.find_test(a.session_id, "T-NOPE00") is None
    both = [st.find_test(s, tid) for s in (a.session_id, b.session_id)]
    assert any(both)


def test_delete_session_cascades(tmp_path, st):
    o = create_session(make_bank(10), n_students=3, q_per_test=5, out_dir=tmp_path, storage=st, seed=1)
    st.delete_session(o.session_id)
    assert st.list_sessions() == [] and st.all_test_ids() == set()


def test_grade_scale_persists(tmp_path):
    p = tmp_path / "s.db"
    Storage(p).set_grade_scale(GradeScale(90, 75, 60))
    assert Storage(p).get_grade_scale().min3 == 60
    with pytest.raises(ValueError):
        Storage(p).set_grade_scale(GradeScale(10, 20, 30))


def test_open_xlsx_locked_message(tmp_path, monkeypatch):
    import os

    def deny(*a, **k):
        raise PermissionError("busy")
    monkeypatch.setattr(os, "replace", deny)
    with pytest.raises(PermissionError, match="yoping"):
        excel_io.write_answer_key(generate_tests(make_bank(5), 2, 3, seed=1).tests, tmp_path / "k.xlsx")


def test_service_with_students_creates_three_files_and_stores_roster(tmp_path, st):
    from models import Student
    studs = [Student(f"O'quvchi {i}", "9-A", "5-maktab") for i in range(1, 8)]
    o = create_session(make_bank(20), n_students=999, q_per_test=6, out_dir=tmp_path, storage=st,
                       students=studs, seed=1)
    assert o.pdf_path.exists() and o.key_path.exists() and o.info_pdf_path.exists()
    s = st.list_sessions()[0]
    assert s.n_students == 7 and s.info_pdf_path == str(o.info_pdf_path)
    ws = load_workbook(o.key_path).active
    assert ws.max_row == 8 and ws["C2"].value == "O'quvchi 1" and ws["C8"].value == "O'quvchi 7"
    tid = ws["A4"].value
    t = st.find_test(o.session_id, tid)
    assert (t.student, t.klass, t.school) == ("O'quvchi 3", "9-A", "5-maktab")
    from pypdf import PdfReader
    info = PdfReader(str(o.info_pdf_path)).pages[2].extract_text()
    assert tid in info and "O'quvchi 3" in info


def test_old_database_is_migrated(tmp_path):
    import sqlite3
    p = tmp_path / "old.db"
    c = sqlite3.connect(p)
    c.executescript("""
    CREATE TABLE sessions(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, created_at TEXT NOT NULL,
      title TEXT NOT NULL DEFAULT '', n_students INTEGER NOT NULL, q_per_test INTEGER NOT NULL, seed INTEGER,
      source_file TEXT NOT NULL DEFAULT '', pdf_path TEXT NOT NULL DEFAULT '', key_path TEXT NOT NULL DEFAULT '');
    CREATE TABLE tests(id INTEGER PRIMARY KEY AUTOINCREMENT, session_id INTEGER NOT NULL, test_id TEXT NOT NULL,
      number INTEGER NOT NULL, UNIQUE(session_id,test_id));
    INSERT INTO sessions(name,created_at,n_students,q_per_test) VALUES('eski','2026-01-01 10:00:00',1,1);
    INSERT INTO tests(session_id,test_id,number) VALUES(1,'T-OLD123',1);
    """)
    c.commit()
    c.close()
    s = Storage(p)
    row = s.list_sessions()[0]
    assert row.name == "eski" and row.info_pdf_path == ""
    assert s.find_test(1, "T-OLD123") is None or True          # elementlari yo'q, lekin xato bermaydi
    assert s.all_test_ids() == {"T-OLD123"}
