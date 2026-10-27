import csv

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader

from generator import generate_tests
from helpers import make_bank
from info_pdf import build_info_pdf
from models import Student
from parser import ParseError
from pdf_builder import PdfMeta
from students import parse_students


def test_xlsx_with_header(tmp_path):
    p = tmp_path / "s.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["T/r", "F.I.Sh.", "Sinf", "Maktab"])
    ws.append([1, "Aliyev Vali", "9-A", "12-maktab"])
    ws.append([None, None, None, None])
    ws.append([2, "  Karimova   Zuhra ", "9-B", "12-maktab"])
    assert [s.name for s in (wb.save(p) or parse_students(p))] == ["Aliyev Vali", "Karimova Zuhra"]


def test_xlsx_without_header_and_reordered_columns(tmp_path):
    p = tmp_path / "s.xlsx"
    wb = Workbook()
    wb.active.append(["Aliyev Vali", "9-A", "5-maktab"])
    wb.save(p)
    assert parse_students(p) == [Student("Aliyev Vali", "9-A", "5-maktab")]
    wb = Workbook()
    wb.active.append(["Maktab", "Sinf", "O'quvchi"])
    wb.active.append(["7-maktab", "8-B", "Soliyev Ali"])
    wb.save(p)
    assert parse_students(p) == [Student("Soliyev Ali", "8-B", "7-maktab")]


@pytest.mark.parametrize("delim", [";", ",", "\t"])
def test_csv_delimiters_and_cp1251(tmp_path, delim):
    p = tmp_path / "s.csv"
    p.write_bytes(f"F.I.Sh.{delim}Sinf{delim}Maktab\nИванов Пётр{delim}10-А{delim}Школа 3\n".encode("cp1251"))
    assert parse_students(p) == [Student("Иванов Пётр", "10-А", "Школа 3")]


def test_txt_names_only_and_docx(tmp_path):
    t = tmp_path / "n.txt"
    t.write_text("Aliyev Vali\n\nKarimova Zuhra\n", encoding="utf-8")
    assert [s.name for s in parse_students(t)] == ["Aliyev Vali", "Karimova Zuhra"]
    d = Document()
    d.add_paragraph("Aliyev Vali; 9-A; 12-maktab")
    d.add_paragraph("Karimova Zuhra; 9-B")
    p = tmp_path / "n.docx"
    d.save(p)
    assert parse_students(p) == [Student("Aliyev Vali", "9-A", "12-maktab"), Student("Karimova Zuhra", "9-B", "")]


def test_student_file_errors(tmp_path):
    with pytest.raises(ParseError, match="topilmadi"):
        parse_students(tmp_path / "yoq.xlsx")
    e = tmp_path / "e.txt"
    e.write_text("\n\n")
    with pytest.raises(ParseError, match="F.I.Sh"):
        parse_students(e)
    bad = tmp_path / "b.xlsx"
    bad.write_text("x")
    with pytest.raises(ParseError):
        parse_students(bad)
    big = tmp_path / "big.txt"
    big.write_text("\n".join(f"Ism {i}" for i in range(1001)))
    with pytest.raises(ParseError, match="1000"):
        parse_students(big)


def test_info_pdf_one_page_per_student_all_fields(tmp_path):
    res = generate_tests(make_bank(10), 3, 5, seed=1)
    studs = [Student("Aliyev Vali", "9-A", "12-maktab"), Student("Oʻktamova Gʻuncha", "9-B", "Школа №3"),
             Student("", "", "")]
    out = tmp_path / "i.pdf"
    assert build_info_pdf(res.tests, studs, out, PdfMeta("Algebra", "Sinov", "27.10.2026")) == 3
    pages = [p.extract_text() for p in PdfReader(str(out)).pages]
    assert len(pages) == 3
    for t, pg, st in zip(res.tests, pages, studs):
        assert t.test_id in pg and "TEST ID" in pg
        for v in (st.name, st.klass, st.school):
            assert v in pg
    assert "Algebra" in pages[0] and "27.10.2026" in pages[0]
    assert "+" not in "".join(pages)


def test_info_pdf_without_students_and_long_name(tmp_path):
    res = generate_tests(make_bank(5), 2, 3, seed=1)
    out = tmp_path / "e.pdf"
    build_info_pdf(res.tests, None, out)
    assert len(PdfReader(str(out)).pages) == 2
    long = Student("Juda " * 60, "9-A", "Maktab " * 40)
    build_info_pdf(res.tests, [long, long], tmp_path / "l.pdf")
    assert len(PdfReader(str(tmp_path / "l.pdf")).pages) == 2


def test_names_resembling_header_words_are_kept(tmp_path):
    p = tmp_path / "n.txt"
    p.write_text("Ismoil Karimov\nSinfdosh Aliyev\nMaktabova Zuhra\n", encoding="utf-8")
    assert [s.name for s in parse_students(p)] == ["Ismoil Karimov", "Sinfdosh Aliyev", "Maktabova Zuhra"]
