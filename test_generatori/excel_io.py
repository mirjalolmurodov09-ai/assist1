"""Excel fayllarga yozish: javoblar kaliti va tekshirish natijalari."""
from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import GeneratedTest

_HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _put(ws, row: int, col: int, value, *, bold: bool = False, fill=None, center: bool = False):
    """Yacheykaga matnni formulaga aylanib ketmasdan yozadi (=, +, -, @ bilan boshlansa ham)."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
    if bold:
        cell.font = Font(bold=True)
    if fill:
        cell.fill = fill
    cell.border = _BORDER
    if center:
        cell.alignment = Alignment(horizontal="center", vertical="center")
    return cell


def _save(wb: Workbook, path: Path) -> None:
    """Faylni xavfsiz saqlaydi (ochiq turgan fayl uchun tushunarli xabar)."""
    tmp = path.with_name(path.name + ".part")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as e:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise PermissionError(
            f"«{path.name}» fayliga yozib bo'lmadi. Faylni yoping va qayta urinib ko'ring.") from e


def write_answer_key(tests: Sequence[GeneratedTest], path: str | Path) -> None:
    """javoblar_kaliti.xlsx: Test ID | Tartib raqam | O'quvchi | 1..Q savollarning to'g'ri javobi."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kalit"
    q = max((len(t.items) for t in tests), default=0)
    heads = ["Test ID", "Tartib raqam", "O'quvchi"] + [f"{i}-savol" for i in range(1, q + 1)]
    for c, h in enumerate(heads, start=1):
        _put(ws, 1, c, h, bold=True, fill=_HEAD_FILL, center=True)
    for r, t in enumerate(tests, start=2):
        _put(ws, r, 1, t.test_id, bold=True)
        _put(ws, r, 2, t.number, center=True)
        _put(ws, r, 3, "")
        for i, L in enumerate(t.correct_letters()):
            _put(ws, r, 4 + i, L, center=True)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 28
    for c in range(4, 4 + q):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.freeze_panes = "D2"
    _save(wb, path)


def write_results(rows: Iterable[dict], path: str | Path) -> None:
    """natijalar.xlsx: Test ID, F.I.Sh., to'g'ri javoblar soni, foiz, baho ..."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    heads = ["Test ID", "Tartib raqam", "F.I.Sh.", "To'g'ri", "Noto'g'ri", "Javobsiz",
             "Savollar soni", "Foiz (%)", "Baho", "Tekshirilgan vaqt"]
    for c, h in enumerate(heads, start=1):
        _put(ws, 1, c, h, bold=True, fill=_HEAD_FILL, center=True)
    r = 1
    for r, row in enumerate(rows, start=2):
        _put(ws, r, 1, row["test_id"], bold=True)
        _put(ws, r, 2, row.get("number", ""), center=True)
        _put(ws, r, 3, row.get("student", "") or "")
        _put(ws, r, 4, row["correct"], center=True)
        _put(ws, r, 5, row["wrong"], center=True)
        _put(ws, r, 6, row["blank"], center=True)
        _put(ws, r, 7, row["total"], center=True)
        _put(ws, r, 8, round(row["percent"], 1), center=True)
        _put(ws, r, 9, row["grade"], center=True)
        _put(ws, r, 10, row.get("checked_at", ""), center=True)
    widths = [14, 14, 30, 9, 11, 10, 14, 10, 7, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    _save(wb, path)
